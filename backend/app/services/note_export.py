"""Render project notes as Markdown, PDF, or Word documents."""
from io import BytesIO
import json
import re
from typing import Iterable
from urllib.parse import quote

import fitz
from docx import Document as WordDocument

from app.db.models import Note

_CONTENT_TYPES = {
    "markdown": "text/markdown; charset=utf-8",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
_INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\x00-\x1f\x7f]+')


def _safe_filename(name: str, format: str) -> str:
    """Keep user-supplied CJK names while removing path and header controls."""
    safe_name = _INVALID_FILENAME_CHARS.sub("_", name).strip(" ._")[:120] or "notes"
    return f"{safe_name}.{format if format != 'markdown' else 'md'}"


def content_disposition(filename: str) -> str:
    """Build a safe attachment header for a possibly non-ASCII filename.

    Args:
        filename: Sanitized filename to send to the browser.

    Returns:
        An RFC 5987 Content-Disposition value.
    """
    return "attachment; filename*=UTF-8''%s" % quote(filename, safe="")


def _note_markdown(note: Note) -> str:
    content = (note.content or "").strip()
    if not content and note.payload_json:
        content = "```json\n%s\n```" % json.dumps(
            note.payload_json, ensure_ascii=False, indent=2
        )
    return "# %s\n\n%s" % (note.title, content)


def _markdown_blocks(markdown: str) -> Iterable[tuple[str, int, str]]:
    for line in markdown.splitlines():
        if not line.strip() or re.fullmatch(r"\s*(?:-{3,}|\*{3,}|_{3,})\s*", line):
            yield "space", 0, ""
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            yield "heading", len(heading.group(1)), heading.group(2)
            continue
        bullet = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if bullet:
            yield "bullet", len(bullet.group(1)) // 2, bullet.group(2)
            continue
        number = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
        if number:
            yield "number", len(number.group(1)) // 2, "%s. %s" % (
                number.group(2), number.group(3)
            )
            continue
        quote = re.match(r"^\s*>\s?(.*)$", line)
        if quote:
            yield "quote", 0, quote.group(1)
            continue
        yield "paragraph", 0, line


def _docx(notes: list[Note]) -> bytes:
    document = WordDocument()
    for note in notes:
        for kind, level, text in _markdown_blocks(_note_markdown(note)):
            if kind == "space":
                continue
            if kind == "heading":
                document.add_heading(text, level=min(level, 6))
            elif kind == "bullet":
                document.add_paragraph(text, style="List Bullet")
            elif kind == "number":
                document.add_paragraph(
                    re.sub(r"^\d+\.\s*", "", text), style="List Number"
                )
            elif kind == "quote":
                document.add_paragraph(text, style="Quote")
            else:
                document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _wrap_pdf_line(
    text: str,
    cjk_font: fitz.Font,
    latin_font: fitz.Font,
    size: float,
    width: float,
) -> list[str]:
    lines, current, current_width = [], "", 0.0
    for char in text:
        font = latin_font if char.isascii() else cjk_font
        char_width = font.text_length(char, size)
        if current and current_width + char_width > width:
            lines.append(current)
            current = char
            current_width = char_width
        else:
            current += char
            current_width += char_width
    if current or not lines:
        lines.append(current)
    return lines


def _insert_pdf_line(
    page: fitz.Page,
    text: str,
    position: tuple[float, float],
    size: float,
    cjk_font: fitz.Font,
    latin_font: fitz.Font,
) -> None:
    x, y = position
    runs = re.findall(r"[^\x00-\x7f]+|[\x00-\x7f]+", text)
    for run in runs:
        font_name, font = (
            ("helv", latin_font) if run.isascii() else ("china-t", cjk_font)
        )
        page.insert_text((x, y), run, fontname=font_name, fontsize=size)
        x += font.text_length(run, size)


def _pdf(notes: list[Note]) -> bytes:
    document = fitz.open()
    font_name = "china-t"
    cjk_font = fitz.Font(fontname=font_name)
    latin_font = fitz.Font(fontname="helv")
    margin = 48
    page = document.new_page()
    y = margin

    for note in notes:
        for kind, level, text in _markdown_blocks(_note_markdown(note)):
            if kind == "space":
                y += 7
                continue
            size = max(11, 21 - 2 * (level - 1)) if kind == "heading" else 11
            indent = (16 * (level + 1)) if kind == "bullet" else 16 if kind in {
                "number", "quote"
            } else 0
            if kind == "bullet":
                text = "- " + text
            elif kind == "quote":
                text = "> " + text
            line_height = size * 1.45
            for line in _wrap_pdf_line(
                text, cjk_font, latin_font, size,
                page.rect.width - 2 * margin - indent,
            ):
                if y + line_height > page.rect.height - margin:
                    page = document.new_page()
                    y = margin
                _insert_pdf_line(
                    page,
                    line,
                    (margin + indent, y + size),
                    size,
                    cjk_font,
                    latin_font,
                )
                y += line_height
            y += 5 if kind == "heading" else 2

    output = document.tobytes()
    document.close()
    return output


def export_notes(
    notes: list[Note], format: str, filename: str
) -> tuple[bytes, str, str]:
    """Render notes in the requested download format.

    Args:
        notes: Notes already filtered to a project the caller can access.
        format: One of ``markdown``, ``pdf``, or ``docx``.
        filename: User-facing note or project title for the download name.

    Returns:
        Export bytes, content type, and safe download filename.
    """
    if format not in _CONTENT_TYPES:
        raise ValueError("Unsupported note export format")
    if format == "markdown":
        content = "\n\n---\n\n".join(_note_markdown(note) for note in notes)
        body = (content.rstrip() + "\n").encode("utf-8")
    elif format == "pdf":
        body = _pdf(notes)
    else:
        body = _docx(notes)
    return body, _CONTENT_TYPES[format], _safe_filename(filename, format)
