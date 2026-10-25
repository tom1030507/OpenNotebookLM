"""Render an existing video summary as a PowerPoint slideshow."""
from io import BytesIO
from typing import List

from pptx import Presentation
from pptx.util import Inches

from app.schemas import VideoSummaryPptxScene


def render_video_summary_pptx(scenes: List[VideoSummaryPptxScene]) -> bytes:
    """Create a 16:9 PowerPoint with one narrated scene on each slide.

    Args:
        scenes: The validated script already shown to the user.

    Returns:
        The presentation as an in-memory PPTX archive.
    """
    presentation = Presentation()
    presentation.slide_width = Inches(40 / 3)
    presentation.slide_height = Inches(7.5)

    for scene in scenes:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = scene.headline
        content = slide.placeholders[1].text_frame
        content.clear()
        for index, bullet in enumerate(scene.bullets):
            paragraph = content.paragraphs[0] if index == 0 else content.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
        slide.notes_slide.notes_text_frame.text = scene.narration

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
