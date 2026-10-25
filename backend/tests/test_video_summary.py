"""Tests for the video summary API.

The service is injected rather than constructed by the router, so these tests
never reach for a model server. With `LLM_MODE` defaulting to `auto` a
module-level service would point at a local provider and every request here
would open a socket to it.
"""
import inspect
from io import BytesIO

import pytest
from fastapi import FastAPI
from pptx import Presentation
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.db.database import get_db
from app.db.models import Base, Chunk, Document, Project, ProjectDocument
from app.routers import video
from app.services.source_digest import project_documents
from app.services.video_summary import VideoSummaryService, VideoSummaryTooLargeError
from conftest import OfflineLLM, auth_headers, authenticated_client, owner_id

app = FastAPI()
app.include_router(video.router, prefix="/api")

engine = create_engine(
    "sqlite:///:memory:",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def override_get_db():
    """Override database dependency for testing."""
    try:
        db = TestingSessionLocal()
        yield db
    finally:
        db.close()


app.dependency_overrides[get_db] = override_get_db
# The script is extracted rather than written here; the generated path is
# covered in `tests/unit/test_video_summary_service.py`.
app.dependency_overrides[video.get_video_summary_service] = (
    lambda: VideoSummaryService(llm_service=OfflineLLM())
)

Base.metadata.create_all(bind=engine)

with TestingSessionLocal() as setup_session:
    client = authenticated_client(app, setup_session)


@pytest.fixture
def two_projects_with_own_documents():
    """Seed two projects, each owning one document with one chunk."""
    with TestingSessionLocal() as db:
        owner = owner_id(db)
        db.add_all([
            Project(id="project-a", user_id=owner, name="Project A"),
            Project(id="project-b", user_id=owner, name="Project B"),
            Document(
                id="document-a",
                user_id=owner,
                title="Only In A",
                source_type="url",
                status="ready",
            ),
            Document(
                id="document-b",
                user_id=owner,
                title="Only In B",
                source_type="url",
                status="ready",
            ),
        ])
        db.flush()
        db.add_all([
            ProjectDocument(project_id="project-a", document_id="document-a"),
            ProjectDocument(project_id="project-b", document_id="document-b"),
            Chunk(
                id="chunk-a",
                document_id="document-a",
                text="Anything",
                heading_path="Only In A/Rainfall",
            ),
            Chunk(
                id="chunk-b",
                document_id="document-b",
                text="Anything",
                heading_path="Only In B/Glaciers",
            ),
        ])
        db.commit()

    yield

    with TestingSessionLocal() as db:
        db.query(Chunk).delete()
        db.query(ProjectDocument).delete()
        db.query(Document).delete()
        db.query(Project).delete()
        db.commit()


@pytest.fixture
def empty_project():
    """Seed a project with no sources attached."""
    with TestingSessionLocal() as db:
        db.add(Project(id="project-empty", user_id=owner_id(db), name="Nothing yet"))
        db.commit()

    yield

    with TestingSessionLocal() as db:
        db.query(Project).delete()
        db.commit()


def test_the_script_opens_on_the_project(two_projects_with_own_documents):
    """The title card is the only place the whole project is stated."""
    response = client.post("/api/projects/project-a/video-summary")

    assert response.status_code == 200
    body = response.json()
    assert body["project_id"] == "project-a"
    assert body["scenes"][0]["kind"] == "title"
    assert body["scenes"][0]["headline"] == "Project A"


def test_each_source_gets_a_scene(two_projects_with_own_documents):
    """One scene per source, and only this project's sources."""
    body = client.post("/api/projects/project-a/video-summary").json()

    sources = [scene for scene in body["scenes"] if scene["kind"] == "source"]
    assert [scene["source_label"] for scene in sources] == ["Only In A"]
    assert sources[0]["document_id"] == "document-a"


def test_another_projects_sources_are_absent(two_projects_with_own_documents):
    """A summary must not leak a sibling project's sources."""
    assert "Only In B" not in client.post(
        "/api/projects/project-a/video-summary",
    ).text


def test_headings_become_the_slide(two_projects_with_own_documents):
    """Without a model, the document's own structure is what is on screen."""
    body = client.post("/api/projects/project-a/video-summary").json()

    assert body["scenes"][1]["bullets"] == ["Rainfall"]


def test_the_script_closes_on_a_recap(two_projects_with_own_documents):
    """The last scene is a recap, composed from the scenes before it."""
    body = client.post("/api/projects/project-a/video-summary").json()

    assert body["scenes"][-1]["kind"] == "closing"
    assert body["scenes"][-1]["bullets"] == ["Only In A"]


def test_the_existing_script_exports_as_a_powerpoint_without_regeneration(
    two_projects_with_own_documents,
):
    """The slides and speaker notes come from the scene script the dialog has."""
    with TestingSessionLocal() as db:
        project = db.query(Project).filter(Project.id == "project-a").one()
        project.name = "../Project A?."
        db.commit()

    scenes = [
        {
            "id": "title",
            "kind": "title",
            "headline": "A prepared title",
            "bullets": ["Prepared bullet"],
            "narration": "Prepared narration.",
            "document_id": None,
            "source_label": None,
        },
        {
            "id": "closing",
            "kind": "closing",
            "headline": "A prepared close",
            "bullets": [],
            "narration": "Prepared closing narration.",
            "document_id": None,
            "source_label": None,
        },
    ]

    response = client.post(
        "/api/projects/project-a/video-summary/pptx",
        json={"scenes": scenes},
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert response.headers["content-disposition"] == (
        "attachment; filename*=utf-8''_Project%20A_.pptx"
    )
    presentation = Presentation(BytesIO(response.content))
    assert len(presentation.slides) == 2
    first_slide = presentation.slides[0]
    assert first_slide.shapes.title.text == "A prepared title"
    assert first_slide.placeholders[1].text == "Prepared bullet"
    assert "Prepared narration." in first_slide.notes_slide.notes_text_frame.text
    assert presentation.slides[1].shapes.title.text == "A prepared close"


def test_powerpoint_export_rejects_malformed_scene_bullets():
    """The file endpoint accepts only the bounded scene shape we generate."""
    scenes = [
        {
            "id": "title",
            "kind": "title",
            "headline": "Title",
            "bullets": ["one", "two", "three", "four", "five"],
            "narration": "Narration.",
        },
        {
            "id": "closing",
            "kind": "closing",
            "headline": "Close",
            "bullets": [],
            "narration": "Closing.",
        },
    ]

    assert client.post(
        "/api/projects/project-a/video-summary/pptx",
        json={"scenes": scenes},
    ).status_code == 422


def test_every_scene_has_narration(two_projects_with_own_documents):
    """The scene advances when the voice finishes, so silence would stall it."""
    body = client.post("/api/projects/project-a/video-summary").json()

    assert all(scene["narration"].strip() for scene in body["scenes"])


def test_scene_count_and_duration_describe_the_script(
    two_projects_with_own_documents,
):
    """Title, one source, closing — and long enough to be worth playing."""
    body = client.post("/api/projects/project-a/video-summary").json()

    assert body["scene_count"] == 3
    assert body["scene_count"] == len(body["scenes"])
    assert body["estimated_seconds"] > 0


def test_oversized_video_summary_returns_payload_too_large(
    two_projects_with_own_documents,
    monkeypatch,
):
    """The UI never receives a script that cannot be exported."""
    def reject_summary(*_args, **_kwargs):
        raise VideoSummaryTooLargeError("Video summary exceeds the export limits")

    monkeypatch.setattr(VideoSummaryService, "generate", reject_summary)

    response = client.post("/api/projects/project-a/video-summary")

    assert response.status_code == 413
    assert response.json()["detail"] == "Video summary exceeds the export limits"


def test_project_document_loader_applies_a_query_limit(two_projects_with_own_documents):
    """The summary can detect oversized projects without loading every source."""
    with TestingSessionLocal() as db:
        project = db.query(Project).filter_by(id="project-a").one()

        documents = project_documents(db, project, limit=1)

    assert len(documents) == 1


def test_the_model_is_reported(two_projects_with_own_documents):
    """`model_used` distinguishes a written script from an extracted one."""
    body = client.post("/api/projects/project-a/video-summary").json()

    assert body["model_used"] == "fallback"


def test_generated_at_carries_a_utc_designator(two_projects_with_own_documents):
    """A designator-less timestamp is read as local time by the browser."""
    body = client.post("/api/projects/project-a/video-summary").json()

    assert body["generated_at"].endswith("Z") or body["generated_at"].endswith("+00:00")


def test_an_empty_project_still_answers(empty_project):
    """A project with no sources answers, rather than failing."""
    response = client.post("/api/projects/project-empty/video-summary")

    assert response.status_code == 200
    body = response.json()
    assert [scene["kind"] for scene in body["scenes"]] == ["title", "closing"]


def test_a_missing_project_is_not_found():
    """An id that names nothing answers 404, not 500."""
    assert client.post("/api/projects/nope/video-summary").status_code == 404


def test_an_anonymous_caller_is_refused():
    """Every data-bearing route sits behind the token check."""
    anonymous = client.__class__(app)

    assert anonymous.post("/api/projects/project-a/video-summary").status_code == 401
    assert client.__class__(
        app, headers=auth_headers(),
    ).post("/api/projects/project-a/video-summary").status_code in (200, 404)


def test_get_cannot_generate_a_video_summary(two_projects_with_own_documents):
    """Reading a URL must not run the project generation command."""
    assert client.get("/api/projects/project-a/video-summary").status_code == 405


def test_the_route_is_not_a_coroutine():
    """A blocking route has to run in the threadpool, not on the event loop.

    Writing a script calls the LLM, which blocks for as long as the model
    takes. The container runs a single uvicorn worker, so an `async def` here
    that never awaits stalls every other request for that whole time —
    `/healthz` included, which the compose healthcheck gives ten seconds.
    FastAPI hands a plain `def` handler to its threadpool, which is where that
    work belongs.
    """
    assert not inspect.iscoroutinefunction(video.project_video_summary)
