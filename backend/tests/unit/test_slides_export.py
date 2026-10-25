from io import BytesIO

from pptx import Presentation

from app.schemas import VideoSummaryPptxScene
from app.services.slides_export import render_video_summary_pptx


def test_rendered_slides_keep_scene_text_and_notes():
    scenes = [
        VideoSummaryPptxScene(
            id="title",
            kind="title",
            headline="Rainfall research",
            bullets=["Rainfall increased"],
            narration="The project reviews rainfall.",
            document_id=None,
            source_label=None,
        ),
        VideoSummaryPptxScene(
            id="closing",
            kind="closing",
            headline="Key point",
            bullets=["Rainfall is rising"],
            narration="That is the main finding.",
            document_id=None,
            source_label=None,
        ),
    ]

    presentation = Presentation(BytesIO(render_video_summary_pptx(scenes)))

    assert len(presentation.slides) == 2
    assert presentation.slides[0].shapes.title.text == "Rainfall research"
    assert presentation.slides[0].placeholders[1].text == "Rainfall increased"
    assert "The project reviews rainfall." in (
        presentation.slides[0].notes_slide.notes_text_frame.text
    )
    assert presentation.slides[1].shapes.title.text == "Key point"
    assert presentation.slide_width / presentation.slide_height == 16 / 9
