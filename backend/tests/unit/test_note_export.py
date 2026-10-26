from io import BytesIO
from types import SimpleNamespace

import fitz
from docx import Document as WordDocument
from docx.shared import Inches

from app.services.note_export import export_notes


def make_note(
    title="研究筆記",
    content=(
        "# 結果\n\n- 水 [1]\n\n1. 步驟 [2]\n\n"
        "**Sources**\n\n- 資料來源"
    ),
):
    return SimpleNamespace(title=title, content=content, payload_json=None)


def test_markdown_export_keeps_headings_lists_and_citations():
    body, content_type, filename = export_notes(
        [make_note()], "markdown", "研究筆記"
    )

    assert content_type == "text/markdown; charset=utf-8"
    assert filename == "研究筆記.md"
    assert "# 研究筆記" in body.decode("utf-8")
    assert "- 水 [1]" in body.decode("utf-8")
    assert "**Sources**" in body.decode("utf-8")


def test_docx_export_keeps_headings_lists_and_citations():
    body, content_type, _ = export_notes([make_note()], "docx", "研究筆記")
    document = WordDocument(BytesIO(body))

    assert content_type == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert [
        (paragraph.text, paragraph.style.name) for paragraph in document.paragraphs
    ] == [
        ("研究筆記", "Heading 1"),
        ("結果", "Heading 1"),
        ("水 [1]", "List Bullet"),
        ("1. 步驟 [2]", "Normal"),
        ("**Sources**", "Normal"),
        ("資料來源", "List Bullet"),
    ]


def test_docx_export_preserves_nested_source_and_excerpt_depth():
    body, _, _ = export_notes(
        [make_note(content="- Source\n  - [1], page 12\n    > Excerpt")],
        "docx",
        "研究筆記",
    )
    paragraphs = WordDocument(BytesIO(body)).paragraphs
    source, citation, excerpt = paragraphs[1:]

    assert [source.text, citation.text, excerpt.text] == [
        "Source", "[1], page 12", "Excerpt",
    ]
    assert source.paragraph_format.left_indent == Inches(0.25)
    assert citation.paragraph_format.left_indent == Inches(0.5)
    assert excerpt.paragraph_format.left_indent == Inches(0.75)


def test_docx_export_preserves_nested_numbering_and_start_values():
    body, _, _ = export_notes(
        [make_note(content="5. Parent\n  1) Child\n6. Next")],
        "docx",
        "研究筆記",
    )
    paragraphs = WordDocument(BytesIO(body)).paragraphs[1:]

    assert [paragraph.text for paragraph in paragraphs] == [
        "5. Parent", "1) Child", "6. Next",
    ]
    assert [paragraph.paragraph_format.left_indent for paragraph in paragraphs] == [
        Inches(0.5), Inches(0.75), Inches(0.5),
    ]
    assert [paragraph.paragraph_format.first_line_indent for paragraph in paragraphs] == [
        Inches(-0.25), Inches(-0.25), Inches(-0.25),
    ]


def test_pdf_export_embeds_cjk_text_and_keeps_citations():
    body, content_type, filename = export_notes([make_note()], "pdf", "研究筆記")
    document = fitz.open(stream=body, filetype="pdf")
    text = "\n".join(page.get_text() for page in document)

    assert content_type == "application/pdf"
    assert filename == "研究筆記.pdf"
    assert "研究筆記" in text
    assert "結果" in text
    assert "水 [1]" in text
    assert "資料來源" in text


def test_pdf_export_indents_nested_citations_and_numbered_lists():
    note = make_note(content=(
        "- Source\n  - [1], page 12\n    > Excerpt\n  1. Next"
    ))
    body, _, _ = export_notes([note], "pdf", note.title)
    document = fitz.open(stream=body, filetype="pdf")
    page = document[0]
    spans = [
        span
        for block in page.get_text("dict")["blocks"]
        for line in block.get("lines", [])
        for span in line["spans"]
    ]
    left = {span["text"]: span["bbox"][0] for span in spans}

    assert left["- Source"] < left["- [1], page 12"]
    assert left["- [1], page 12"] < left["> Excerpt"]
    assert left["- Source"] < left["1. Next"]


def test_export_filename_keeps_cjk_and_removes_path_and_header_characters():
    _, _, filename = export_notes(
        [make_note("研究/筆記\r\n")], "markdown", "研究/筆記\r\n"
    )

    assert filename == "研究_筆記.md"
