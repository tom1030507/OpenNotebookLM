"""Tests for local text and Office source extraction."""
from pathlib import Path
import zipfile

import pytest

from app.adapters.office import (
    InvalidOfficeFileError,
    OfficeAdapter,
)


def write_docx(path: Path, document_xml: str | None = None) -> None:
    """Create a small OpenXML Word document without storing a binary fixture."""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Override PartName="/word/document.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            '</Types>',
        )
        archive.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/></Relationships>',
        )
        archive.writestr(
            "word/document.xml",
            document_xml or (
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:body><w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>'
                '<w:r><w:t>Chapter One</w:t></w:r></w:p>'
                '<w:p><w:r><w:t>Useful paragraph content.</w:t></w:r></w:p>'
                '<w:sectPr/></w:body></w:document>'
            ),
        )


def write_epub(path: Path) -> None:
    """Create a minimal EPUB 3 package with one chapter."""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr(
            "META-INF/container.xml",
            '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf" '
            'media-type="application/oebps-package+xml"/></rootfiles></container>',
        )
        archive.writestr(
            "OEBPS/content.opf",
            '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book-id">'
            '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
            '<dc:identifier id="book-id">book</dc:identifier><dc:title>Sample Book</dc:title>'
            '<dc:language>en</dc:language></metadata><manifest>'
            '<item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/>'
            '</manifest><spine><itemref idref="chapter"/></spine></package>',
        )
        archive.writestr(
            "OEBPS/chapter.xhtml",
            '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Chapter</title></head>'
            '<body><h1>Book Chapter</h1><p>EPUB chapter content.</p></body></html>',
        )


def write_pptx(path: Path) -> None:
    """Create a minimal two-slide package without third-party image libraries.

    Args:
        path: Destination path for the generated presentation.

    Returns:
        None.
    """
    presentation_namespace = (
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    )
    drawing_namespace = (
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    )

    def slide_xml(title: str) -> str:
        """Return one slide containing a title and body placeholder."""
        return (
            '<p:sld %s><p:cSld><p:spTree>'
            '<p:sp><p:nvSpPr><p:cNvPr id="2" name="Title"/><p:cNvSpPr/>'
            '<p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>'
            '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>%s</a:t></a:r></a:p></p:txBody>'
            '</p:sp><p:sp><p:nvSpPr><p:cNvPr id="3" name="Body"/><p:cNvSpPr/>'
            '<p:nvPr><p:ph type="body"/></p:nvPr></p:nvSpPr>'
            '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>Slide body text</a:t></a:r></a:p></p:txBody>'
            '</p:sp></p:spTree></p:cSld></p:sld>'
            % (presentation_namespace + " " + drawing_namespace, title)
        )

    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'
            '<Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
            '<Override PartName="/ppt/slides/slide2.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
            '</Types>',
        )
        archive.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>'
            '</Relationships>',
        )
        archive.writestr(
            "ppt/presentation.xml",
            '<p:presentation %s><p:sldIdLst><p:sldId id="256" r:id="rId1"/>'
            '<p:sldId id="257" r:id="rId2"/></p:sldIdLst></p:presentation>'
            % presentation_namespace,
        )
        archive.writestr(
            "ppt/_rels/presentation.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide2.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>'
            '</Relationships>',
        )
        for index, title in enumerate(("Findings", "Opening"), start=1):
            archive.writestr(
                "ppt/slides/slide%s.xml" % index,
                slide_xml(title),
            )


def rewrite_zip_members(
    path: Path,
    replacements: dict[str, bytes],
    added: dict[str, bytes] | None = None,
) -> None:
    """Replace selected ZIP members without changing the others.

    Args:
        path: ZIP archive to rewrite in place.
        replacements: Existing member names and replacement bytes.
        added: Optional new member names and bytes.

    Returns:
        None.
    """
    with zipfile.ZipFile(path) as source:
        members = [
            (part, replacements.get(part.filename, source.read(part.filename)))
            for part in source.infolist()
        ]

    with zipfile.ZipFile(path, "w") as target:
        for part, data in members:
            target.writestr(part, data)
        for name, data in (added or {}).items():
            target.writestr(name, data)


def result_for(path: Path) -> dict:
    """Extract a generated fixture through the production adapter."""
    return OfficeAdapter().extract_text_from_file(str(path))


def test_text_and_markdown_are_utf8_decoded_with_replacement(tmp_path):
    """Plain text input remains local and malformed UTF-8 stays readable."""
    text_path = tmp_path / "notes.txt"
    text_path.write_bytes(b"first line\n\xffsecond line")
    markdown_path = tmp_path / "guide.md"
    markdown_path.write_text("# Guide\n\nA useful section.", encoding="utf-8")

    text_result = result_for(text_path)
    markdown_result = result_for(markdown_path)

    assert "first line\n\ufffdsecond line" == text_result["text"]
    assert "# Guide" in markdown_result["text"]
    assert text_result["pages"] == []
    assert text_result["num_pages"] == 0
    assert set(text_result) == {"text", "pages", "num_pages", "metadata"}


def test_markdown_frontmatter_and_wikilinks_are_metadata(tmp_path):
    path = tmp_path / "note.md"
    path.write_text(
        "---\ntitle: 資料庫筆記\ntags:\n  - Obsidian\n---\n"
        "# Fallback title\n\nSee [[資料庫|database]] and [[索引#設計]].\n",
        encoding="utf-8",
    )

    result = result_for(path)

    assert result["text"] == "# Fallback title\n\nSee [[資料庫|database]] and [[索引#設計]]."
    assert result["metadata"]["title"] == "資料庫筆記"
    assert result["metadata"]["frontmatter"] == {
        "title": "資料庫筆記",
        "tags": ["Obsidian"],
    }
    assert result["metadata"]["wikilinks"] == ["資料庫", "索引"]


def test_malformed_markdown_frontmatter_stays_in_source_text(tmp_path):
    path = tmp_path / "broken.md"
    path.write_text("---\ntitle: [unterminated\n---\nBody", encoding="utf-8")

    result = result_for(path)

    assert result["text"].startswith("---\ntitle: [unterminated")
    assert "frontmatter" not in result["metadata"]


def test_source_code_files_are_read_as_utf8(tmp_path):
    path = tmp_path / "main.py"
    path.write_text("print('資料庫')", encoding="utf-8")

    result = result_for(path)

    assert result["text"] == "print('資料庫')"
    assert result["metadata"]["format"] == "py"


def test_docx_heading_style_becomes_markdown_heading(tmp_path):
    """Word heading styles survive conversion for section citations."""
    path = tmp_path / "report.docx"
    write_docx(path)

    result = result_for(path)

    assert "# Chapter One" in result["text"]
    assert "Useful paragraph content." in result["text"]


def test_pptx_has_one_markdown_heading_per_slide(tmp_path):
    """Each slide title is a heading that can become a citation section."""
    path = tmp_path / "deck.pptx"
    write_pptx(path)

    text = result_for(path)["text"]

    assert "## Slide 1: Opening" in text
    assert "## Slide 2: Findings" in text
    assert text.index("## Slide 1: Opening") < text.index("## Slide 2: Findings")
    assert text.count("Slide body text") == 2
    assert text.count("## Slide 1:") == 1
    assert text.count("## Slide 2:") == 1


def test_pptx_with_unresolved_slide_relationship_is_rejected(tmp_path):
    """Presentation slide references must resolve before enqueue."""
    path = tmp_path / "missing-slide-relationship.pptx"
    write_pptx(path)
    with zipfile.ZipFile(path) as archive:
        presentation = archive.read("ppt/presentation.xml")
    rewrite_zip_members(
        path,
        {
            "ppt/presentation.xml": presentation.replace(
                b'r:id="rId1"',
                b'r:id="rId9"',
                1,
            ),
        },
    )

    with pytest.raises(InvalidOfficeFileError, match="slide relationship"):
        OfficeAdapter().validate_file(str(path), ".pptx")


def test_pptx_with_wrong_slide_root_is_rejected(tmp_path):
    """A relationship target must contain a presentation slide part."""
    path = tmp_path / "wrong-slide-root.pptx"
    write_pptx(path)
    with zipfile.ZipFile(path) as archive:
        slide = archive.read("ppt/slides/slide1.xml")
    rewrite_zip_members(
        path,
        {
            "ppt/slides/slide1.xml": slide
            .replace(b"<p:sld ", b"<p:other ", 1)
            .replace(b"</p:sld>", b"</p:other>", 1),
        },
    )

    with pytest.raises(InvalidOfficeFileError, match="slide part"):
        OfficeAdapter().validate_file(str(path), ".pptx")


def test_xlsx_has_sheet_headings_and_markdown_tables(tmp_path):
    """Each workbook sheet is named and represented as a Markdown table."""
    from openpyxl import Workbook

    path = tmp_path / "metrics.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quarterly Data"
    sheet.append(["Region", "Revenue"])
    sheet.append(["North", 42])
    workbook.create_sheet("Notes").append(["Comment"])
    workbook.save(path)

    text = result_for(path)["text"]

    assert "## Sheet: Quarterly Data" in text
    assert "## Sheet: Notes" in text
    assert "| Region | Revenue |" in text
    assert "| North | 42 |" in text


def test_epub_chapter_headings_are_preserved(tmp_path):
    """EPUB chapters keep their headings when converted to searchable text."""
    path = tmp_path / "book.epub"
    write_epub(path)

    text = result_for(path)["text"]

    assert "Book Chapter" in text
    assert "EPUB chapter content." in text


def test_non_archive_renamed_as_docx_is_rejected(tmp_path):
    """A renamed executable is refused before it can enter the job queue."""
    path = tmp_path / "renamed.exe.docx"
    path.write_bytes(b"MZ\x00\x01not a Word document")

    with pytest.raises(InvalidOfficeFileError, match="valid Office document"):
        OfficeAdapter().validate_file(str(path), ".docx")


def test_docx_with_malformed_required_xml_is_rejected(tmp_path):
    """A package with the right names but broken XML cannot enter the queue."""
    path = tmp_path / "malformed.docx"
    write_docx(path, document_xml="<w:document><w:body>")

    with pytest.raises(InvalidOfficeFileError, match="valid Office document"):
        OfficeAdapter().validate_file(str(path), ".docx")


def test_docx_with_well_formed_wrong_root_is_rejected(tmp_path):
    """A valid XML document in the main part is not necessarily a DOCX part."""
    path = tmp_path / "wrong-root.docx"
    write_docx(
        path,
        document_xml=(
            '<w:notDocument xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
        ),
    )

    with pytest.raises(InvalidOfficeFileError, match="valid Office document"):
        OfficeAdapter().validate_file(str(path), ".docx")


def test_docx_with_missing_relationship_target_is_rejected(tmp_path):
    """Internal relationship targets must exist in the package."""
    path = tmp_path / "missing-target.docx"
    write_docx(path)
    with zipfile.ZipFile(path) as archive:
        relationships = archive.read("_rels/.rels")
    rewrite_zip_members(
        path,
        {
            "_rels/.rels": relationships.replace(
                b"word/document.xml",
                b"word/missing.xml",
            ),
        },
    )

    with pytest.raises(InvalidOfficeFileError, match="missing relationship target"):
        OfficeAdapter().validate_file(str(path), ".docx")


def test_docx_with_wrong_main_relationship_target_is_rejected(tmp_path):
    """The package root must identify the actual Word document part."""
    path = tmp_path / "wrong-main-target.docx"
    write_docx(path)
    with zipfile.ZipFile(path) as archive:
        relationships = archive.read("_rels/.rels")
    rewrite_zip_members(
        path,
        {
            "_rels/.rels": relationships.replace(
                b"word/document.xml",
                b"word/other.xml",
            ),
        },
        added={"word/other.xml": b"<other/>"},
    )

    with pytest.raises(InvalidOfficeFileError, match="main-part relationship"):
        OfficeAdapter().validate_file(str(path), ".docx")


def test_epub_without_rootfile_is_rejected(tmp_path):
    """An EPUB container must declare a package document for conversion."""
    path = tmp_path / "no-rootfile.epub"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "mimetype",
            "application/epub+zip",
            compress_type=zipfile.ZIP_STORED,
        )
        archive.writestr(
            "META-INF/container.xml",
            '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0"/>',
        )

    with pytest.raises(InvalidOfficeFileError, match="rootfile"):
        OfficeAdapter().validate_file(str(path), ".epub")


def test_docx_with_corrupted_package_part_is_rejected(tmp_path):
    """A bad CRC in a non-XML member is caught before MarkItDown reads it."""
    path = tmp_path / "corrupted.docx"
    write_docx(path)
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("word/media/image1.png", b"image bytes")
    with zipfile.ZipFile(path) as archive:
        part = archive.getinfo("word/media/image1.png")
    data_offset = (
        part.header_offset
        + 30
        + len(part.filename.encode("utf-8"))
        + len(part.extra)
    )
    with path.open("r+b") as package:
        package.seek(data_offset)
        package.write(bytes([package.read(1)[0] ^ 1]))

    with pytest.raises(InvalidOfficeFileError, match="corrupted package part"):
        OfficeAdapter().validate_file(str(path), ".docx")


def test_xlsx_above_cell_cap_is_rejected(tmp_path, monkeypatch):
    """The configured worksheet cell ceiling is checked before pandas loads it."""
    from openpyxl import Workbook
    import app.adapters.office as office

    path = tmp_path / "too-many-cells.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["one", "two", "three"])
    workbook.save(path)
    monkeypatch.setattr(office, "MAX_SPREADSHEET_CELLS", 2)

    with pytest.raises(InvalidOfficeFileError, match="cell limit"):
        OfficeAdapter().validate_file(str(path), ".xlsx")


def test_xlsx_within_the_cell_cap_is_accepted(tmp_path):
    """A small valid spreadsheet passes content validation."""
    from openpyxl import Workbook

    path = tmp_path / "small.xlsx"
    workbook = Workbook()
    workbook.active.append(["one", "two"])
    workbook.save(path)

    OfficeAdapter().validate_file(str(path), ".xlsx")


def test_xlsx_sparse_cells_exceeding_grid_cap_are_rejected(tmp_path, monkeypatch):
    """Far-apart populated cells cannot expand into an oversized DataFrame."""
    from openpyxl import Workbook
    import app.adapters.office as office

    path = tmp_path / "sparse.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "first"
    sheet["Z100"] = "last"
    workbook.save(path)
    monkeypatch.setattr(office, "MAX_SPREADSHEET_CELLS", 10)

    with pytest.raises(InvalidOfficeFileError, match="cell limit"):
        OfficeAdapter().validate_file(str(path), ".xlsx")


def test_xlsx_relocated_worksheet_is_included_in_cell_cap(tmp_path, monkeypatch):
    """Workbook relationship targets are checked even outside the usual folder."""
    from openpyxl import Workbook
    import app.adapters.office as office

    path = tmp_path / "relocated.xlsx"
    workbook = Workbook()
    workbook.active.title = "Empty"
    populated = workbook.create_sheet("Sparse")
    populated["A1"] = "first"
    populated["Z100"] = "last"
    workbook.save(path)

    with zipfile.ZipFile(path) as archive:
        relationships = archive.read("xl/_rels/workbook.xml.rels")
        content_types = archive.read("[Content_Types].xml")
        populated_sheet = archive.read("xl/worksheets/sheet2.xml")
    empty_sheet = (
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b"<sheetData/></worksheet>"
    )
    rewrite_zip_members(
        path,
        {
            "xl/_rels/workbook.xml.rels": relationships.replace(
                b"/xl/worksheets/sheet2.xml",
                b"/xl/custom.xml",
            ),
            "[Content_Types].xml": content_types.replace(
                b"/xl/worksheets/sheet2.xml",
                b"/xl/custom.xml",
            ),
            "xl/worksheets/sheet2.xml": empty_sheet,
        },
        added={"xl/custom.xml": populated_sheet},
    )
    monkeypatch.setattr(office, "MAX_SPREADSHEET_CELLS", 10)

    with pytest.raises(InvalidOfficeFileError, match="cell limit"):
        OfficeAdapter().validate_file(str(path), ".xlsx")


def test_xlsx_row_index_cannot_hide_a_large_row_span(tmp_path, monkeypatch):
    """The row parser's index matters even when cells claim a nearby row."""
    from openpyxl import Workbook
    import app.adapters.office as office

    path = tmp_path / "large-row-index.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "value"
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        worksheet = archive.read("xl/worksheets/sheet1.xml")
    assert b'<row r="1"' in worksheet
    rewrite_zip_members(
        path,
        {
            "xl/worksheets/sheet1.xml": worksheet.replace(
                b'<row r="1"',
                b'<row r="1048576"',
                1,
            ),
        },
    )
    monkeypatch.setattr(office, "MAX_SPREADSHEET_CELLS", 10)

    with pytest.raises(InvalidOfficeFileError, match="cell limit|row reference"):
        OfficeAdapter().validate_file(str(path), ".xlsx")


def test_xlsx_missing_row_index_uses_the_sequential_row(tmp_path, monkeypatch):
    """Optional row indices can be inferred without losing the row bound."""
    from openpyxl import Workbook
    import app.adapters.office as office

    path = tmp_path / "implicit-row-index.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "first"
    workbook.active["A2"] = "second"
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        worksheet = archive.read("xl/worksheets/sheet1.xml")
    assert b'<row r="2"' in worksheet
    rewrite_zip_members(
        path,
        {
            "xl/worksheets/sheet1.xml": worksheet.replace(
                b'<row r="2"',
                b"<row",
                1,
            ),
        },
    )
    monkeypatch.setattr(office, "MAX_SPREADSHEET_CELLS", 2)

    OfficeAdapter().validate_file(str(path), ".xlsx")


def test_xlsx_out_of_range_row_index_is_rejected(tmp_path):
    """Rows beyond Excel's maximum cannot bypass the conversion bound."""
    from openpyxl import Workbook

    path = tmp_path / "out-of-range-row-index.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "value"
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        worksheet = archive.read("xl/worksheets/sheet1.xml")
    rewrite_zip_members(
        path,
        {
            "xl/worksheets/sheet1.xml": worksheet.replace(
                b'<row r="1"',
                b'<row r="1048577"',
                1,
            ),
        },
    )

    with pytest.raises(InvalidOfficeFileError, match="invalid row references"):
        OfficeAdapter().validate_file(str(path), ".xlsx")
